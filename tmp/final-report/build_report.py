from pathlib import Path
from copy import deepcopy
import hashlib, json, re
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT=Path(__file__).resolve().parents[2]
TMP=Path(__file__).parent
REFERENCE=ROOT/'Docs/TH1_IOT_V5.docx'
OUT=ROOT/'Docs/Report/BAO_CAO_BTL_IOT_HOANG_QUOC_ANH.docx'
EV=ROOT/'Docs/Evidence/2026-10-27-final-review'
OLD=ROOT/'Docs/Evidence/2026-10-21'
FIG=TMP/'figures'
doc=Document(REFERENCE)
# Retain the original cover, logo, section geometry, styles, relationships and footer.
# The rest is a new concise report, explicitly replacing the longer TH1 body.
body=doc.element.body
for el in list(body)[17:]:
 if el.tag!=qn('w:sectPr'):body.remove(el)
for p in doc.paragraphs:
 if 'BÁO CÁO BÀI THỰC HÀNH 1' in p.text:
  p.text='BÁO CÁO BÀI TẬP LỚN'
  p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  for r in p.runs:r.bold=True;r.font.name='Times New Roman';r.font.size=Pt(15)
 if p.text.startswith('Đề tài:'):
  if 'Title' not in doc.styles:doc.styles.add_style('Title',WD_STYLE_TYPE.PARAGRAPH)
  p.style='Title'
  p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(18);r.bold=True;r.font.color.rgb=RGBColor(0,0,0)
  pp=p._p.get_or_add_pPr()
  for item in list(pp):
   if item.tag==qn('w:pBdr'):pp.remove(item)

for name,size,bold in [('Normal',12,False),('Title',18,True),('Heading 1',16,True),('Heading 2',14,True),('Heading 3',12,True),('Caption',11,False),('TOC 1',12,True),('TOC 2',11,False)]:
 if name not in doc.styles:doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
 s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.bold=bold;s.font.color.rgb=RGBColor(0,0,0)
 s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
 f=s.paragraph_format;f.line_spacing=1.12;f.space_after=Pt(5);f.space_before=Pt(0)
 if name.startswith('Heading'):
  f.space_before=Pt(8);f.space_after=Pt(6);f.keep_with_next=True
 if name.startswith('TOC'):f.space_after=Pt(3);f.space_before=Pt(0)
 if name=='Caption':f.space_after=Pt(7);f.space_before=Pt(3)
 # Remove theme font/color and paragraph rules inherited from the reference.
 rp=s._element.get_or_add_rPr()
 for c in rp.findall(qn('w:color')):
  for a in ('themeColor','themeTint','themeShade'):
   c.attrib.pop(qn('w:'+a),None)
 for c in rp.findall(qn('w:rFonts')):
  for a in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
   c.attrib.pop(qn('w:'+a),None)
 pp=s._element.get_or_add_pPr()
 for x in list(pp):
  if x.tag==qn('w:pBdr'):pp.remove(x)

normal=doc.styles['Normal'].paragraph_format
normal.first_line_indent=Cm(0);normal.widow_control=True
chapter=0;tnum=0
INDEX_ENTRIES={'F':[],'T':[]}
INDEX_PLACEHOLDERS={}
def page():doc.add_page_break()
def h(s,level=2):
 p=doc.add_paragraph(s,style=f'Heading {level}');p.paragraph_format.keep_with_next=True
 return p
def chapter_title(n,title):
 global chapter,tnum
 chapter=n;tnum=0
 return h(f'CHƯƠNG {n} {title.upper()}',1)
def p(s='',bold=False,italic=False,align=None,size=None):
 q=doc.add_paragraph();q.alignment=align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
 r=q.add_run(s);r.bold=bold;r.italic=italic
 if size:r.font.size=Pt(size)
 return q
def bullet(s):
 q=p('•  '+s);q.paragraph_format.left_indent=Cm(.25);q.paragraph_format.first_line_indent=Cm(-.25);return q
def caption(s):
 q=doc.add_paragraph(s,style='Caption');q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.keep_with_next=False
 index_entry(q,s,'F')
 return q
def index_entry(q,s,kind):
 n=len(INDEX_ENTRIES[kind])+1;name=f'IOT_{kind}_{n}';bid=str(90000+(0 if kind=='F' else 1000)+n)
 start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),bid);start.set(qn('w:name'),name);q._p.insert(0,start)
 end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),bid);q._p.append(end)
 INDEX_ENTRIES[kind].append((s,name))
def add_index(title,kind):
 p(title,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,size=16)
 INDEX_PLACEHOLDERS[kind]=doc.add_paragraph()
def populate_indexes():
 for kind,holder in INDEX_PLACEHOLDERS.items():
  for label,bookmark in INDEX_ENTRIES[kind]:
   q=Paragraph(OxmlElement('w:p'),holder._parent);holder._p.addprevious(q._p)
   q.paragraph_format.space_after=Pt(3);q.paragraph_format.line_spacing=1.1
   q.paragraph_format.tab_stops.add_tab_stop(Cm(16.5),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
   r=q.add_run(label+'\t');r.font.size=Pt(11)
   r=q.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'begin');r._r.append(a)
   r=q.add_run();i=OxmlElement('w:instrText');i.set(qn('xml:space'),'preserve');i.text=f' PAGEREF {bookmark} \\h ';r._r.append(i)
   a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'separate');r._r.append(a)
   q.add_run('0').font.size=Pt(11)
   a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'end');q.add_run()._r.append(a)
  holder._p.getparent().remove(holder._p)
def picture(path,label,width=16.5,crop=None):
 path=Path(path); iw,ih=Image.open(path).size
 l,t,r,b=crop or (0,0,0,0)
 cw=iw*(1-l-r);ch=ih*(1-t-b)
 q=doc.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_after=Pt(1);q.paragraph_format.keep_with_next=True
 sh=q.add_run().add_picture(str(path),width=Cm(width),height=Cm(width*ch/cw))
 sh._inline.docPr.set('descr',label)
 if crop:
  fill=sh._inline.graphic.graphicData.pic.blipFill
  src=OxmlElement('a:srcRect')
  for k,v in zip(('l','t','r','b'),crop):src.set(k,str(round(v*100000)))
  fill.insert(1,src)
 caption(label)
def table(headers,rows,widths,caption_text=None,font=10.5):
 global tnum
 if caption_text:
  tnum+=1;cap=doc.add_paragraph(f'Bảng {chapter}.{tnum}  {caption_text}',style='Caption');cap.paragraph_format.keep_with_next=True
  index_entry(cap,cap.text,'T')
 tb=doc.add_table(rows=1, cols=len(headers));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
 for col,w in zip(tb.columns,widths):col.width=Cm(w)
 for i,word in enumerate(headers):tb.rows[0].cells[i].text=word
 for row in rows:
  cells=tb.add_row().cells
  for i,val in enumerate(row):cells[i].text=str(val)
 for ri,row in enumerate(tb.rows):
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if ri==0:trpr.append(OxmlElement('w:tblHeader'))
  for ci,(cell,w) in enumerate(zip(row.cells,widths)):
   cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for side in ('top','left','bottom','right'):
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'65');el.set(qn('w:type'),'dxa');mar.append(el)
   pr.append(mar);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E5E7EB' if ri==0 else 'FFFFFF');pr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for side in ('top','left','bottom','right'):
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   pr.append(borders)
   for q in cell.paragraphs:
    q.paragraph_format.space_after=Pt(1);q.paragraph_format.space_before=Pt(1);q.paragraph_format.line_spacing=1.02;q.paragraph_format.keep_with_next=False
    if ci==0 and widths[ci]<2:q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    else:q.alignment=WD_ALIGN_PARAGRAPH.LEFT
    for run in q.runs:run.font.name='Times New Roman';run.font.size=Pt(font);run.bold=ri==0;run.font.color.rgb=RGBColor(0,0,0)
 spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(3);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=Pt(1);spacer.add_run().font.size=Pt(1)
 return tb
def code(s):
 q=doc.add_paragraph();q.paragraph_format.space_after=Pt(4)
 r=q.add_run(s);r.font.name='Consolas';r.font.size=Pt(9.5)
 return q

page()
p('MỤC LỤC',bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,size=16)
q=doc.add_paragraph();r=q.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'begin');r._r.append(a)
i=OxmlElement('w:instrText');i.set(qn('xml:space'),'preserve');i.text=' TOC \\o "1-2" \\h \\z \\u ';r._r.append(i)
sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r._r.append(sep)
r=q.add_run('Mục lục được cập nhật theo tiêu đề trong tài liệu.')
end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');q.add_run()._r.append(end)

page();add_index('DANH MỤC HÌNH','F')
page();add_index('DANH MỤC BẢNG','T')
p('TỪ VIẾT TẮT',bold=True,size=12,align=WD_ALIGN_PARAGRAPH.LEFT)
p('IoT: Internet of Things; MQTT: Message Queuing Telemetry Transport; REST: Representational State Transfer; JWT: JSON Web Token; STOMP: Simple Text Oriented Messaging Protocol; GPIO: General Purpose Input Output; ADC: Analog to Digital Converter; ERD: Entity Relationship Diagram; UTC: Coordinated Universal Time.',size=10.5,align=WD_ALIGN_PARAGRAPH.LEFT)

page();chapter_title(1,'Giới thiệu')
h('1.1 Bài toán và mục tiêu')
p('Đề tài xây dựng hệ thống Web theo dõi nhiệt độ, độ ẩm, ánh sáng và điều khiển hai LED từ ESP32. Số đo được lưu vào MySQL, hiển thị realtime và có thể tra cứu. Mỗi lệnh điều khiển được ghi người thực hiện, theo dõi xác nhận và thông báo khi không có phản hồi.')
table(['Mục tiêu','Kết quả cần đạt'],[
 ('Thu thập','Đọc DHT11 và LDR; truyền số đo qua MQTT.'),
 ('Giám sát','Hiển thị số đo cuối, một biểu đồ ba đường và trạng thái phần cứng.'),
 ('Điều khiển','ON/OFF hai LED; chỉ cập nhật trạng thái theo xác nhận hợp lệ.'),
 ('Tra cứu','Tìm kiếm, lọc, sắp xếp và phân trang cảm biến/lịch sử.'),
 ('Người dùng','Đăng nhập bằng JWT; xem/sửa hồ sơ và ghi username cho từng lệnh.'),
 ('Triển khai','Khởi động dịch vụ bằng Docker; cung cấp Swagger/OpenAPI.')],[3.6,12.9],'Mục tiêu đề tài')
h('1.2 Phạm vi thực hiện')
p('Phần cứng gồm ESP32, DHT11, LDR và hai LED trong mạng Wi-Fi nội bộ. Ứng dụng có năm màn hình: Đăng nhập, Dashboard, Dữ liệu cảm biến, Lịch sử điều khiển và Hồ sơ. Ba tài khoản user123, user124, user125 có cùng nhóm chức năng; không có đăng ký hoặc phân quyền nhiều vai trò.')
p('Lux được ước lượng từ ADC, chưa hiệu chuẩn. LED là thiết bị chấp hành của mô hình; phản hồi MQTT xác nhận mức GPIO theo firmware. Hệ thống chưa điều khiển thiết bị điện lưới.')
h('1.3 Phương pháp và bố cục')
p('Quá trình thực hiện gồm phân tích từ TH1, thiết kế dữ liệu/giao diện, xây dựng ứng dụng, tích hợp ESP32 và kiểm tra xuyên suốt MQTT–Backend–MySQL–WebSocket–Web. Báo cáo gồm bốn chương: Giới thiệu, Phân tích, Thiết kế và Kết luận. Kết quả kiểm thử sử dụng lượt chốt ngày 27/10/2026 [1], [2].')

page();chapter_title(2,'Phân tích')
h('2.1 Tác nhân và chức năng')
p('User sử dụng Web; ESP32 gửi số đo và phản hồi trạng thái. Các chức năng sau đăng nhập dùng chung giữa ba tài khoản, riêng hồ sơ được truy xuất theo chủ thể JWT.')
picture(FIG/'usecase.png','Hình 2.1  Các use case và tác nhân của hệ thống',width=16.5)
table(['Mã','Chức năng','Kết quả'],[
 ('UC01','Đăng nhập','Xác thực tài khoản và cấp JWT.'),('UC02','Giám sát Dashboard','Số đo cuối, biểu đồ và tình trạng phần cứng.'),
 ('UC03','Điều khiển LED','Gửi ON/OFF; chờ xác nhận hoặc timeout.'),('UC04','Tra cứu cảm biến','Tìm theo đại lượng/giá trị/thời gian, sort/page.'),
 ('UC05','Tra cứu điều khiển','Lọc lệnh và xem đúng username người gửi.'),('UC06','Quản lý hồ sơ','Xem, sửa, lưu hoặc hủy thay đổi.')],[1.7,4.5,10.3],'Yêu cầu chức năng')

page();h('2.2 Dữ liệu và quy tắc nghiệp vụ')
table(['Mã nguồn đo','Đại lượng','Đơn vị'],[
 ('DHT11_TEMP','Nhiệt độ từ DHT11','°C'),('DHT11_HUM','Độ ẩm từ DHT11','%RH'),('LDR_LIGHT','Ánh sáng quy đổi từ ADC','lux ước lượng')],[5.4,7.6,3.5],'Danh mục cảm biến')
table(['Quy tắc','Nội dung'],[
 ('Bản ghi','Một reading hợp lệ tạo một SensorData; không có sample gộp ba cảm biến.'),
 ('Biểu đồ','Tối đa 15 reading mới nhất mỗi Sensor. X thể hiện thứ tự, nhãn giờ theo nhiệt độ; tooltip giữ giờ riêng từng Sensor.'),
 ('Kết nối','Không có telemetry hợp lệ trong 30 giây thì Offline; giữ số đo cuối.'),
 ('Lệnh LED','Mỗi lần bấm tạo ActionHistory, kể cả lệnh trùng trạng thái. requestId = ActionHistory.id.'),
 ('Xác nhận','202 chỉ là tiếp nhận. Chỉ status MQTT hợp lệ mới xác nhận lịch sử và cập nhật Device.'),
 ('Timeout','Sau 10 giây chưa có xác nhận: TIMEOUT, Device giữ trạng thái đã xác nhận gần nhất.'),
 ('Processing','Hiển thị khi deliveryState=PENDING; cột action/status trong SQL vẫn chỉ ON/OFF.'),
 ('Người gửi','Lấy User.id từ JWT; hiển thị User.username qua quan hệ lịch sử, không nhận actor từ body.'),
 ('Truy vấn','UI mặc định 5 dòng; cho chọn 5/10/15/20. Filter trước sort và page.'),
 ('Thời gian','Lưu và truyền JSON theo UTC; giao diện hiển thị giờ Việt Nam.')],[3.1,13.4],'Quy tắc chính')
h('2.3 Yêu cầu phi chức năng')
p('API nghiệp vụ cần JWT; mật khẩu lưu BCrypt. Sự kiện realtime được phát sau khi dữ liệu commit. Giao diện thể hiện trạng thái tải/rỗng/lỗi/chờ, không tự tạo số đo hoặc giả lập ACK. Sau kết nối lại, Frontend tải snapshot REST để phục hồi trạng thái; cấu hình và secrets được tách khỏi source [2].')

page();chapter_title(3,'Thiết kế')
h('3.1 Kiến trúc và môi trường triển khai')
picture(FIG/'architecture.png','Hình 3.1  Kiến trúc và chiều trao đổi dữ liệu')
p('MQTT kết nối phần cứng với Backend theo mô hình publish/subscribe [3]. REST tải snapshot, tra cứu và tiếp nhận thao tác. STOMP trên WebSocket đẩy thay đổi tới React [4]. Frontend không truy cập trực tiếp MySQL hoặc broker.')
table(['Thành phần','Triển khai và kết nối'],[
 ('React + Nginx + Spring Boot','Container app; Web tại localhost:8088. Nginx proxy /api và /ws tới Backend:8080.'),
 ('Mosquitto','Service trong Compose. Backend: mosquitto:1883; ESP32: IP LAN laptop:1325.'),
 ('MySQL','Container sẵn có, volume được giữ. Backend: host.docker.internal:3307; MySQL nội bộ:3306.'),
 ('ESP32','Kết nối Wi-Fi và MQTT thật; không dùng localhost hoặc tên service Docker.')],[4.3,12.2],'Môi trường hiện tại')
p('Compose quản lý app và broker; MySQL hiện là container riêng. start.ps1 khởi động MySQL hiện có trước Compose. Chế độ BTH4 tắt IOT_LED_DEMO_ENABLED; không có responder mô phỏng trạng thái LED.')

page();h('3.2 Phần cứng và firmware')
picture(FIG/'wiring.png','Hình 3.2  Ánh xạ GPIO theo sketch của dự án',width=15.5)
table(['Tín hiệu','GPIO','Chức năng'],[('DHT11 DATA','4','Đọc nhiệt độ và độ ẩm.'),('LDR AO','34','Đọc ADC analog.'),('LED1','18','Điều khiển LED 1 qua điện trở.'),('LED2','19','Điều khiển LED 2 qua điện trở.')],[4.1,2,10.4],'Kết nối thiết kế')
p('Firmware khởi tạo GPIO và cảm biến, kết nối Wi-Fi/MQTT rồi lặp xử lý lệnh và đọc cảm biến theo cấu hình khoảng 2 giây. LDR lấy trung bình 10 mẫu ADC; hàm adcToLux quy đổi xấp xỉ 0–200 lux. analogRead cung cấp giá trị ADC, không tự tạo phép đo lux đã hiệu chuẩn [7].')
p('Khi nhận command, ESP32 kiểm tra requestId/deviceCode/action, đổi GPIO và publish status cùng requestId. Các gói của ba đại lượng được gửi riêng; thời gian Backend nhận có thể khác nhau.')
p('Chèn ảnh mạch toàn cảnh và cận chân nối tại đây.',italic=True,align=WD_ALIGN_PARAGRAPH.CENTER,size=11)
caption('Hình 3.3  Mạch thực tế của hệ thống')
p('Chèn ảnh Serial Monitor thể hiện số đo và phản hồi LED tại đây.',italic=True,align=WD_ALIGN_PARAGRAPH.CENTER,size=11)
caption('Hình 3.4  Số đo và phản hồi trên Serial Monitor')

page();h('3.3 Cơ sở dữ liệu')
picture(FIG/'erd.png','Hình 3.5  ERD rút gọn gồm năm bảng',width=15.7)
table(['Quan hệ','Ý nghĩa'],[
 ('Sensor 1–N SensorData','Một Sensor có nhiều reading.'),('User 1–N ActionHistory','Một User gửi nhiều lệnh.'),('Device 1–N ActionHistory','Một LED có nhiều yêu cầu điều khiển.')],[7,9.5],'Quan hệ dữ liệu')
p('Schema gồm năm bảng, 40 cột. SensorData.value dùng DECIMAL(12,4), thời gian dùng DATETIME(6) theo UTC [6]. Chỉ mục ghép (sensorId, recordedAt) và (deviceId, createdAt) hỗ trợ tra cứu. Từ điển đầy đủ đặt tại Phụ lục.')
p('ActionHistory lưu userId/deviceId/action/status/createdAt/confirmedAt. confirmedAt=NULL khi chưa xác nhận; deliveryState được suy ra theo thời gian. Username trả qua JOIN User, không thêm cột username vào ActionHistory. Seed chỉ thêm tài khoản thiếu, không ghi đè hồ sơ hoặc mật khẩu đã có.')

page();h('3.4 Đăng nhập')
picture(OLD/'images/F08-login-empty-no-credentials.png','Hình 3.6  Giao diện đăng nhập',width=9.3,crop=(.25,.21,.25,.18))
picture(FIG/'login.png','Hình 3.7  Sequence đăng nhập',width=16.5)
table(['API','Dữ liệu và phản hồi'],[('POST /api/auth/login','Body: username, password. Hợp lệ: 200 {token}; sai tài khoản/mật khẩu:401; thiếu trường:400.')],[6,10.5],'API đăng nhập')
p('JWT dùng HS256, issuer iot-backend, subject User.id; REST gửi Authorization: Bearer. Token được kiểm tra chữ ký và hạn dùng; Frontend lưu phiên trong localStorage và chuyển Login khi hết hạn/401. Thời hạn mặc định 3.600 giây, có thể cấu hình [2].')
p('Kiểm tra 27/10: cả ba tài khoản đăng nhập được; mật khẩu sai, token thiếu/sai/hết hạn bị từ chối. UI đăng nhập lại thành công sau khi phiên cũ hết hạn.',size=11)

page();h('3.5 Dashboard và dữ liệu realtime')
picture(EV/'images/dashboard-final-after-reload.png','Hình 3.8  Dashboard Online sau kiểm tra với ESP32',crop=(1/6,0,0,.16))
p('Dashboard tải GET /api/dashboard rồi đăng ký realtime bằng useEffect; cleanup hủy listener khi rời trang [5]. Khi STOMP connected/reconnected, tải lại snapshot. Event sensors cập nhật reading/card; devices cập nhật LED; hardware đổi banner; notifications xử lý timeout.')
table(['Đường','Mã dữ liệu','Đơn vị'],[('Đỏ','DHT11_TEMP','°C'),('Xanh','DHT11_HUM','%RH'),('Vàng','LDR_LIGHT','lux')],[3.2,8.3,5],'Ánh xạ ba đường biểu đồ')
p('Mỗi đường giữ 15 reading mới nhất, X cách đều từ vị trí 1 đến 15. Nhãn giờ lấy từ nhiệt độ; tooltip giữ thời gian thật từng Sensor. Khi đủ dữ liệu, ba đường thẳng hàng hai đầu; không nội suy giá trị hay ghép thành chu kỳ đo chung.')

page();h('Luồng dữ liệu và hợp đồng giao tiếp',3)
picture(FIG/'telemetry.png','Hình 3.9  Sequence từ ESP32 tới Dashboard')
table(['Topic MQTT','JSON','Chiều'],[
 ('iot/sensor/data','sensorCode, value','ESP32 → Backend'),('iot/device/command','requestId, deviceCode, action','Backend → ESP32'),('iot/device/status','requestId, deviceCode, status','ESP32 → Backend')],[5.9,6.3,4.3],'Ba topic MQTT của dự án',font=10)
p('MQTT dùng QoS 0, không retain; Backend bỏ gói telemetry/status retained và payload không hợp lệ. Telemetry chỉ nhận Sensor đã đăng ký, active, giá trị hữu hạn và đúng miền. recordedAt do Backend gán, không phải thời điểm lấy mẫu trên ESP32.')
table(['STOMP destination','Dữ liệu gửi tới Frontend'],[
 ('/topic/sensors','Một SensorData đã commit, gồm id/mã/giá trị/đơn vị/giờ.'),('/topic/hardware','ONLINE/OFFLINE và lastSeenAt.'),('/topic/devices','Trạng thái Device hiện tại và request đã xác nhận.'),('/topic/notifications','DEVICE_TIMEOUT và requestId.')],[6,10.5],'Các kênh realtime',font=10)
p('GET /api/dashboard trả latest, chart, devices và hardware. JWT được gửi ở STOMP CONNECT tại /ws; chỉ cho subscribe bốn topic trên, không nhận client SEND. Lượt chốt: 522 MQTT reading = 522 STOMP = 522 SQL khớp, 174 mỗi Sensor; Dashboard tự Offline/Online khi rút/cắm [2].',size=11)

page();h('3.6 Điều khiển LED có xác nhận')
picture(EV/'images/led-pending-on-off-colors.png','Hình 3.10  Chờ bật màu xanh và chờ tắt màu đỏ',crop=(1/6,.65,0,.13))
picture(EV/'images/history-processing-pending-row.png','Hình 3.11  Request 625 đang Processing khi chưa có xác nhận',width=14.5,crop=(1/6,0,0,.43))
picture(EV/'images/dashboard-timeout-message.png','Hình 3.12  Timeout giữ nguyên trạng thái đã xác nhận',crop=(1/6,.65,0,.13))
table(['deliveryState','Điều kiện','Tác động'],[
 ('PENDING','Chưa xác nhận, dưới 10 giây.','Hiện Processing; giữ Device.status.'),('CONFIRMED','Status hợp lệ đúng requestId/deviceCode.','Ghi confirmedAt, cập nhật theo ACK.'),('TIMEOUT','Chưa ACK sau 10 giây.','Báo lỗi; không tự đổi đèn.')],[3.4,6.1,7],'Vòng đời yêu cầu điều khiển')
p('Thông báo pending ON xanh/OFF đỏ và không có nền; nút cho thao tác tiếp theo sáng hơn. Trạng thái đèn chỉ đổi theo xác nhận, không đổi ngay khi người dùng bấm.',size=11)

page();h('Luồng xử lý và kết quả điều khiển',3)
picture(FIG/'led.png','Hình 3.13  Sequence điều khiển LED có requestId')
p('POST /api/devices/{deviceId}/actions; body chỉ có action=ON hoặc OFF. Backend lấy User từ JWT, lưu history rồi publish command. HTTP 202 trả requestId và status=ACCEPTED; publish thất bại trả 503 kèm requestId, history vẫn được giữ.')
code('{"requestId":617,"deviceCode":"LED1","action":"ON"}\n{"requestId":617,"deviceCode":"LED1","status":"ON"}')
table(['Request','Người gửi và thao tác','Kết quả'],[
 ('617–618','user123; LED1 ON/OFF qua REST.','CONFIRMED'),('619–620','user124; LED2 ON/OFF qua REST.','CONFIRMED'),('621–622','user125; LED1 ON/OFF qua REST.','CONFIRMED'),('623–626','user123; Web khi đã rút ESP32.','TIMEOUT'),('627–628','user123; Web LED2 ON/OFF sau cắm lại.','CONFIRMED')],[2.7,10,3.8],'Các yêu cầu chủ động ngày 27/10',font=10)
p('Tám xác nhận khớp MQTT, STOMP, SQL và REST; bốn timeout không có status, confirmedAt=NULL. ACK lặp không đổi xác nhận đầu; ACK cũ không ghi đè Device đã xác nhận yêu cầu mới hơn. ACK muộn có thể hoàn tất history nhưng chưa được thử với phần cứng ở lượt này [2].',size=11)

page();h('3.7 Dữ liệu cảm biến')
picture(EV/'images/sensor-time-search-final.png','Hình 3.14  Tìm đúng một giây trên dữ liệu thật',crop=(1/6,0,0,.44))
picture(FIG/'sensor-search.png','Hình 3.15  Sequence tra cứu dữ liệu cảm biến')
p('GET /api/sensor-data nhận page, size, sortBy, order, search, sensorType, from/to. Chọn dropdown Tìm theo chỉ chuẩn bị trường tìm kiếm; bấm Tìm kiếm mới áp dụng. Chọn Thời gian sẽ đổi ngày/giây theo giờ Việt Nam thành khoảng UTC.')
p('Response: {content, page, size, totalElements, totalPages}; mỗi hàng gồm id, sensorCode, sensorType, value, unit, recordedAt. Filter trước sort/page. UI cho 5/10/15/20 dòng, số dòng ở trái và tổng số ở phải; đổi thứ tự hoặc điều kiện đưa về trang đầu.')
p('Kiểm tra: đổi dropdown chưa lọc ngay; tìm 27.5 chỉ trả nhiệt độ phù hợp; tìm 27/10/2026 18:44:00 trả ba reading đúng giây. Realtime gom tín hiệu tải lại trong 300 ms, giữ bảng trong lúc chờ [2].',size=11)

page();h('3.8 Lịch sử điều khiển')
picture(EV/'images/history-six-columns-desktop.png','Hình 3.16  Lịch sử có Người thực hiện và đầy đủ ngày giờ',width=15.4,crop=(1/6,0,0,.23))
picture(FIG/'history.png','Hình 3.17  Sequence truy vấn lịch sử và username',width=15.5)
p('GET /api/action-history nhận page/size/sortBy/order, search, device, action, status, from/to. Lọc Processing gửi status=PROCESSING, tương ứng deliveryState=PENDING. Xóa lọc đưa các điều kiện về mặc định.')
p('HistoryRow trả id, deviceId, deviceCode, username, action, status, createdAt, confirmedAt, deliveryState. Người thực hiện lấy từ ActionHistory.userId → User.username; người đang xem không thay tác giả. Hai cột Lệnh/Trạng thái không thay cho thông tin deliveryState.')
p('Kiểm tra: page 5/10/15/20, bộ lọc và username khớp SQL; request 617–622 giữ đúng ba người gửi. Desktop đủ sáu cột, mobile cuộn bảng riêng [2].',size=11)

page();h('3.9 Hồ sơ người dùng')
picture(OLD/'images/F20-profile-view.png','Hình 3.18  Giao diện xem hồ sơ và tài nguyên dự án',width=14,crop=(1/6,.06,0,.2))
picture(FIG/'profile.png','Hình 3.19  Sequence xem và cập nhật hồ sơ')
p('GET /api/profile đọc User theo JWT; PUT /api/profile thay tám trường: fullName, studentCode, email, githubUrl, figmaUrl, apiDocsUrl, reportUrl, avatarUrl. Ba trường đầu bắt buộc; avatar chỉ nhận HTTPS. Không sửa id, username hoặc passwordHash.')
p('Lưu hợp lệ trả 200; dữ liệu sai 400; email/mã sinh viên trùng 409. Hủy bỏ thay đổi chưa lưu. Kiểm tra: PUT tạm user125, GET xác nhận rồi khôi phục tám trường; user123 không đổi. Đã thử email trùng, avatar HTTP và body giả username. updatedAt thay đổi do PUT [2].',size=11)

page();h('3.10 Tài liệu API và vận hành')
picture(EV/'images/swagger-seven-endpoints.png','Hình 3.20  Swagger hiển thị bảy thao tác REST',width=14)
p('Swagger UI tại /swagger-ui/index.html, OpenAPI tại /v3/api-docs. Bảy thao tác gồm Login, Dashboard, SensorData, ActionHistory, gửi action và GET/PUT Profile. Schema có HistoryRow.username. API nghiệp vụ dùng bearer JWT; tài liệu công khai không đồng nghĩa API được bỏ xác thực [8].')
p('Khởi động bằng .\\start.ps1 ở thư mục dự án; script dùng MySQL container hiện có và Compose app/broker. Kiểm tra docker compose ps, health app và Dashboard. Cấu hình DB/MQTT/JWT nằm trong file môi trường được gitignore; không đưa secrets vào báo cáo hoặc Git.')
p('Dừng app/broker bằng docker compose --env-file Backend/.env stop. Không dùng down -v để tránh mất dữ liệu; không chạy thêm Backend cùng nhận telemetry. Mosquitto Windows không tự tranh cổng 1325. Firmware giữ IP LAN phù hợp; nội bộ Docker dùng tên service [2].',size=11)

page();chapter_title(4,'Kết luận')
h('4.1 Kết quả đạt được')
p('Hệ thống đã thực hiện thu thập cảm biến, giám sát realtime, điều khiển LED có xác nhận, tra cứu và quản lý hồ sơ. Ba tài khoản có cùng chức năng, lịch sử lưu đúng người gửi; dữ liệu được quản lý bằng năm bảng và có tài liệu API.')
table(['Hạng mục kiểm tra','Kết quả ngày 27/10/2026'],[
 ('Telemetry','522 MQTT reading, 522 STOMP, 174/Sensor; 522 hàng SQL khớp, 0 sai lệch.'),
 ('LED Online','8 request CONFIRMED: 617–622 và 627–628; đúng thiết bị, action và username.'),
 ('Offline và timeout','4 request 623–626 không ACK; có timeout, Device giữ OFF. Dashboard tự Online sau cắm lại.'),
 ('Giao diện','15 điểm mỗi đường thẳng hàng đầu/cuối; tooltip, chú giải, tìm giá trị/thời gian, Processing, màu ON/OFF và reload hoạt động.'),
 ('API và hồ sơ','142 mục kiểm tra PASS, 0 FAIL, 1 OBSERVED; PUT/khôi phục user125 và các nhánh xác thực đã nêu đạt.'),
 ('Tự động và build','12 test Frontend, 4 test Backend H2/MQTT mock đạt; Vite và Docker app image build thành công.')],[5.1,11.4],'Tổng hợp kết quả')
p('Phiên chốt 18:41:54–18:50:36 ngày 27/10/2026 được lưu tại bộ final-review của dự án. Các số lượng trên chỉ thuộc tập đã đối chiếu; 142 mục kiểm tra không phải 142 endpoint độc lập. Test H2/MQTT mock được tách khỏi kết quả ESP32 thật [2].',size=11)
h('4.2 Hạn chế')
p('Telemetry chưa có sampleId hoặc timestamp đo tại ESP32; không suy ra ba Sensor cùng chu kỳ từ giờ nhận. Lux chưa hiệu chuẩn; ACK chỉ phản ánh GPIO, chưa có cảm biến xác nhận LED phát sáng độc lập. Chưa kiểm thử broker/Wi-Fi lỗi, ACK muộn/lặp/đảo thứ tự với phần cứng thật, tải lớn hoặc chạy dài hạn; môi trường nội bộ chưa có TLS đầy đủ.')
h('4.3 Hướng phát triển')
p('Bổ sung mã chu kỳ và timestamp đo; hiệu chuẩn ánh sáng; kiểm thử reconnect/ACK bất thường và đo từng chặng truyền. Sau đó hoàn thiện TLS, sao lưu/phục hồi và cơ chế chống ghi trùng nếu mở rộng Backend. Các thay đổi MQTT phải đồng bộ firmware, Backend và Frontend.')

page();h('TÀI LIỆU THAM KHẢO',1)
refs=[
 ('[1]','Hoàng Quốc Anh. Bài thực hành 1 Hệ thống IoT giám sát môi trường và điều khiển thiết bị. TH1_IOT_V5.docx, 2026.'),
 ('[2]','Mã nguồn dự án và bộ kiểm tra 2026-10-27-final-review: compose.yaml, Backend/CONTRACT.md, schema.sql, firmware IoTMonitor.ino; api-review.json, telemetry-comparison.json, planned-command-results.json và ảnh giao diện. Chốt ngày 27/10/2026.'),
 ('[3]','MQTT.org. MQTT Specifications. https://mqtt.org/mqtt-specification/'),
 ('[4]','Spring. STOMP over WebSocket. https://docs.spring.io/spring-framework/reference/web/websocket/stomp.html'),
 ('[5]','React. useEffect. https://react.dev/reference/react/useEffect'),
 ('[6]','Oracle. MySQL 8.0 Fixed Point Types. https://dev.mysql.com/doc/refman/8.0/en/fixed-point-types.html'),
 ('[7]','Espressif. Arduino ESP32 ADC API. https://docs.espressif.com/projects/arduino-esp32/en/latest/api/adc.html'),
 ('[8]','Swagger. Bearer Authentication. https://swagger.io/docs/specification/v3_0/authentication/bearer-authentication/')]
for a,b in refs:p(a+'  '+b,size=11,align=WD_ALIGN_PARAGRAPH.LEFT)
p('Các tài liệu trực tuyến được truy cập ngày 27/10/2026.',italic=True,size=10.5)

page();h('PHỤ LỤC TỪ ĐIỂN DỮ LIỆU',1)
p('Năm bảng gồm 40 cột; kiểu và khả năng NULL được đối chiếu MySQL. PK: khóa chính; FK: khóa ngoại; UQ: duy nhất. DATETIME(6) lưu tới microsecond theo quy ước UTC của ứng dụng.',size=11)
md=(ROOT/'Docs/Report/PHU_LUC_DU_LIEU.md').read_text(encoding='utf-8-sig')
for idx,name in enumerate(['User','Sensor','SensorData','Device','ActionHistory']):
 if name=='SensorData':page()
 h(name,3)
 part=md.split('## '+name+'\n',1)[1].split('\n## ',1)[0]
 rows=[]
 for line in part.splitlines():
  if line.startswith('|'):
   cells=[v.strip() for v in line.split('|')[1:-1]]
   if len(cells)==5 and cells[0] not in ['Cột','---']:
    cons=cells[3].replace('UNIQUE','UQ').replace('PK, AUTO_INCREMENT','PK, tự tăng')
    if cells[0]=='sensorId':cons='FK → Sensor.id'
    if name=='ActionHistory' and cells[0]=='userId':cons='FK → User.id'
    if name=='ActionHistory' and cells[0]=='deviceId':cons='FK → Device.id'
    rows.append((cells[0],cells[1],cells[2],cons))
 table(['Cột','Kiểu SQL','NULL','Ràng buộc'],rows,[5.1,4,2,5.4],font=10.5)

populate_indexes()
settings=doc.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
doc.core_properties.title='Hệ thống IoT giám sát môi trường và điều khiển thiết bị'
doc.core_properties.subject='Báo cáo bài tập lớn IoT và ứng dụng'
doc.core_properties.author='Hoàng Quốc Anh'
doc.core_properties.keywords='ESP32 MQTT Spring Boot React MySQL STOMP'
doc.core_properties.comments=''
doc.save(OUT)
print(json.dumps({'output':str(OUT),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'figures':len(doc.inline_shapes),'referenceUnchangedSha256':hashlib.sha256(REFERENCE.read_bytes()).hexdigest()},ensure_ascii=False))
